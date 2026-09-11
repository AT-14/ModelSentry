import argparse
import json
from pathlib import Path
from urllib.parse import urlparse

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import pymupdf
import qrcode
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
EVIDENCE = ROOT / "artifacts" / "validation_extended_v2_holdout"
SUBMISSION = ROOT / "submission"

NAVY = "071018"
PANEL = "0D1C27"
PANEL_LIGHT = "132A38"
INK = "F7FAFC"
MUTED = "9AABB7"
GOLD = "F4B942"
TEAL = "69D4C5"
RED = "EF6A6A"
LINE = "294453"

SLIDE_W = 13.333
SLIDE_H = 7.5

TEAM = "Team TSA"
MEMBERS = (
    "Abdallah Abu Zaid",
    "Nahyan Alali",
    "Mahmoud Kassem",
    "Zayed Almarzooqi",
    "Rashed Alnuaimi",
)
UNIVERSITY = "United Arab Emirates University"
REPOSITORY = "https://github.com/AT-14/ModelSentry"
SOURCE_REVISION = "92a244ee98faa8b1dddf0770849ad4961f7fd79c"
VIDEO_LINK_PENDING = "VIDEO LINK PENDING"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def validate_public_url(value: str) -> str:
    parsed = urlparse(value)
    if parsed.scheme != "https" or not parsed.netloc:
        raise ValueError("The video URL must be a public HTTPS URL")
    return value


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    width: float,
    height: float,
    size: float,
    color: str = INK,
    bold: bool = False,
    font: str = "Aptos",
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin: float = 0.0,
    hyperlink: str | None = None,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    if hyperlink:
        run.hyperlink.address = hyperlink
    run.font.color.rgb = rgb(color)
    return box


def add_panel(slide, x: float, y: float, width: float, height: float, fill: str = PANEL):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(LINE)
    shape.line.width = Pt(1)
    shape.adjustments[0] = 0.08
    return shape


def add_link_overlay(
    slide, x: float, y: float, width: float, height: float, url: str
) -> None:
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height),
    )
    shape.fill.background()
    shape.line.fill.background()
    shape.click_action.hyperlink.address = url


def add_metric(slide, x: float, y: float, width: float, value: str, label: str, accent: str = GOLD):
    add_panel(slide, x, y, width, 1.03)
    add_text(slide, value, x + 0.16, y + 0.17, width - 0.30, 0.40, 22, accent, True, "Aptos Display")
    add_text(slide, label.upper(), x + 0.16, y + 0.66, width - 0.30, 0.22, 8, MUTED, True)


def add_base(slide, kicker: str, title: str, page: int, draft: bool) -> None:
    background = slide.background
    background.fill.solid()
    background.fill.fore_color.rgb = rgb(NAVY)
    add_text(slide, kicker.upper(), 0.55, 0.25, 7.0, 0.25, 8.5, TEAL, True)
    if draft:
        add_text(slide, "DRAFT | VIDEO LINK PENDING", 9.0, 0.25, 3.75, 0.25, 8, RED, True, align=PP_ALIGN.RIGHT)
    add_text(slide, title, 0.55, 0.64, 12.1, 0.64, 25, INK, True, "Aptos Display")
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(0.55),
        Inches(7.15),
        Inches(12.75),
        Inches(7.15),
    )
    line.line.color.rgb = rgb(LINE)
    line.line.width = Pt(0.8)
    add_text(slide, "SCHOOL OF CYBER DEFENCE 2026", 0.55, 7.2, 4.0, 0.18, 6.5, MUTED)
    add_text(slide, f"MODELSENTRY  |  {page}/5", 10.0, 7.2, 2.75, 0.18, 6.5, MUTED, align=PP_ALIGN.RIGHT)


def add_flow(slide, items: tuple[tuple[str, str], ...], y: float) -> None:
    x_positions = (0.55, 3.72, 6.89, 10.06)
    for index, ((heading, detail), x) in enumerate(zip(items, x_positions)):
        add_panel(slide, x, y, 2.35, 1.2)
        add_text(slide, heading, x + 0.17, y + 0.18, 2.0, 0.26, 11, INK, True)
        add_text(slide, detail, x + 0.17, y + 0.56, 2.0, 0.48, 8, MUTED)
        if index < 3:
            arrow = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT,
                Inches(x + 2.45),
                Inches(y + 0.60),
                Inches(x_positions[index + 1] - 0.10),
                Inches(y + 0.60),
            )
            arrow.line.color.rgb = rgb(GOLD)
            arrow.line.width = Pt(1.8)
            arrow.line.end_arrowhead = True


def set_notes(slide, _text: str) -> None:
    slide.notes_slide.notes_text_frame.text = ""


def make_qr(url: str, destination: Path) -> None:
    image = qrcode.make(url)
    image.save(destination)


def make_holdout_chart(summary: dict, destination: Path) -> None:
    modes = ("rate_only", "model_aware", "full", "enhanced")
    labels = ("Rate only", "Model aware", "Baseline full", "Enhanced V2.4")
    overview = summary["mode_overview"]
    detection = [overview[mode]["attack_detection_rate"] * 100 for mode in modes]
    benign = [overview[mode]["benign_false_positive_rate"] * 100 for mode in modes]
    positions = np.arange(len(modes))

    figure, axis = plt.subplots(figsize=(10, 5.2), facecolor=f"#{NAVY}")
    axis.set_facecolor(f"#{PANEL}")
    width = 0.34
    detection_bars = axis.bar(
        positions - width / 2,
        detection,
        width,
        color=f"#{GOLD}",
        label="Attack detection rate",
    )
    benign_bars = axis.bar(
        positions + width / 2,
        benign,
        width,
        color=f"#{TEAL}",
        label="Benign mitigation rate",
    )
    axis.bar_label(
        detection_bars,
        labels=[f"{value:.1f}%" for value in detection],
        padding=3,
        color=f"#{INK}",
        fontsize=9,
    )
    axis.bar_label(
        benign_bars,
        labels=[f"{value:.1f}%" for value in benign],
        padding=3,
        color=f"#{INK}",
        fontsize=9,
    )
    axis.set_ylim(0, 108)
    axis.set_ylabel("Share of evaluated sessions (%)", color=f"#{MUTED}")
    axis.set_xticks(positions, labels, color=f"#{INK}")
    axis.tick_params(axis="y", colors=f"#{MUTED}")
    axis.grid(axis="y", color=f"#{LINE}", linewidth=0.7, alpha=0.7)
    axis.set_axisbelow(True)
    for spine in axis.spines.values():
        spine.set_color(f"#{LINE}")
    legend = axis.legend(frameon=False, loc="upper left")
    for text in legend.get_texts():
        text.set_color(f"#{INK}")
    figure.tight_layout()
    figure.savefig(destination, dpi=180, facecolor=f"#{NAVY}", bbox_inches="tight")
    plt.close(figure)


def make_dashboard_snapshot(summary: dict, destination: Path) -> None:
    enhanced = summary["mode_overview"]["enhanced"]
    latency = summary["latency"]["enhanced/api_in_process"]
    figure, axis = plt.subplots(figsize=(13.333, 7.5), facecolor=f"#{NAVY}")
    axis.set_facecolor(f"#{NAVY}")
    axis.set_xlim(0, 1)
    axis.set_ylim(0, 1)
    axis.axis("off")
    axis.text(
        0.05,
        0.92,
        "MODEL IP DEFENSE CONTROL ROOM",
        color=f"#{TEAL}",
        fontsize=10,
        weight="bold",
    )
    axis.text(0.05, 0.83, "ModelSentry", color=f"#{INK}", fontsize=30, weight="bold")
    axis.text(
        0.05,
        0.77,
        "Frozen V2.4 holdout | three untouched seeds",
        color=f"#{MUTED}",
        fontsize=12,
    )
    cards = (
        (
            f"{enhanced['attack_runs_detected']}/{enhanced['attack_runs_tested']}",
            "ATTACK RUNS DETECTED",
        ),
        (
            f"{enhanced['benign_sessions_mitigated']}/{enhanced['benign_sessions_tested']}",
            "BENIGN SESSIONS MITIGATED",
        ),
        (f"{enhanced['final_fidelity']['mean']:.2%}", "MEAN FINAL FIDELITY"),
        (f"{latency['p50_ms']['mean']:.2f} ms", "API P50 LATENCY"),
    )
    for index, (value, label) in enumerate(cards):
        x = 0.05 + index * 0.232
        card = plt.Rectangle(
            (x, 0.55),
            0.21,
            0.14,
            facecolor=f"#{PANEL}",
            edgecolor=f"#{LINE}",
        )
        axis.add_patch(card)
        accent = GOLD if index % 2 == 0 else TEAL
        axis.text(
            x + 0.015,
            0.625,
            value,
            color=f"#{accent}",
            fontsize=19,
            weight="bold",
        )
        axis.text(
            x + 0.015,
            0.575,
            label,
            color=f"#{MUTED}",
            fontsize=7.5,
            weight="bold",
        )
    scenario_values = (("FAST", 3), ("REPLAY", 3), ("DISTRIBUTED", 3), ("SLOW", 2))
    axis.text(
        0.05,
        0.45,
        "DETECTION BY ATTACK SCENARIO",
        color=f"#{INK}",
        fontsize=11,
        weight="bold",
    )
    for index, (name, detected) in enumerate(scenario_values):
        y = 0.37 - index * 0.07
        axis.text(0.05, y, name, color=f"#{MUTED}", fontsize=9, weight="bold")
        axis.add_patch(
            plt.Rectangle((0.19, y - 0.006), 0.48, 0.025, facecolor=f"#{PANEL_LIGHT}")
        )
        axis.add_patch(
            plt.Rectangle(
                (0.19, y - 0.006),
                0.48 * detected / 3,
                0.025,
                facecolor=f"#{TEAL if detected == 3 else GOLD}",
            )
        )
        axis.text(0.69, y, f"{detected}/3", color=f"#{INK}", fontsize=9, weight="bold")
    axis.text(0.76, 0.43, "HONEST LIMIT", color=f"#{GOLD}", fontsize=9, weight="bold")
    axis.text(0.76, 0.36, "One slow-adaptive run", color=f"#{INK}", fontsize=13, weight="bold")
    axis.text(0.76, 0.31, "was not detected.", color=f"#{INK}", fontsize=13, weight="bold")
    axis.text(0.76, 0.22, "88/88 manifest hashes verified", color=f"#{TEAL}", fontsize=9)
    axis.text(0.05, 0.06, f"SOURCE {SOURCE_REVISION}", color=f"#{MUTED}", fontsize=7)
    figure.savefig(destination, dpi=180, facecolor=f"#{NAVY}", bbox_inches="tight")
    plt.close(figure)


def add_slide_one(prs: Presentation, summary: dict, draft: bool) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_base(slide, "Title slide", "ModelSentry", 1, draft)
    add_text(slide, "Detect model-theft behavior and limit API leakage.", 0.55, 1.38, 10.8, 0.5, 20, GOLD, False, "Aptos Display")
    add_text(
        slide,
        "A stateful API-security layer that identifies systematic model learning, explains the evidence, and limits information exposure as extraction behavior emerges.",
        0.55,
        1.98,
        11.7,
        0.75,
        11,
        INK,
    )
    enhanced = summary["mode_overview"]["enhanced"]
    api_latency = summary["latency"]["enhanced/api_in_process"]["p50_ms"]["mean"]
    add_metric(slide, 0.55, 3.05, 2.85, "11/12", "Attack runs detected")
    add_metric(slide, 3.57, 3.05, 2.85, "0/90", "Benign sessions mitigated", TEAL)
    add_metric(
        slide,
        6.59,
        3.05,
        2.85,
        f"{enhanced['final_fidelity']['mean']:.2%}",
        "Mean final fidelity",
    )
    add_metric(slide, 9.61, 3.05, 2.85, f"{api_latency:.2f} ms", "In-process API p50", TEAL)
    add_panel(slide, 0.55, 4.55, 11.91, 1.65)
    add_text(slide, TEAM, 0.78, 4.82, 2.0, 0.32, 14, INK, True)
    add_text(slide, "  |  ".join(MEMBERS), 0.78, 5.28, 10.9, 0.28, 8.5, INK)
    add_text(slide, UNIVERSITY, 0.78, 5.70, 6.0, 0.30, 10, MUTED)
    add_text(slide, "CONTROLLED PROOF OF CONCEPT", 9.2, 4.85, 2.95, 0.25, 8, TEAL, True, align=PP_ALIGN.RIGHT)
    set_notes(
        slide,
        "[Timing: 0:00-0:35]\nOpen with the risk: a public prediction API can become a labeling service for a competitor. ModelSentry watches behavior rather than isolated requests. State the frozen V2.4 holdout result exactly: 11 of 12 attacks detected, no mitigation in 90 benign sessions, and 71.78 percent mean final surrogate fidelity. Disclose that one slow-adaptive run was missed. Transition: first, show how API access becomes model theft.",
    )


def add_slide_two(prs: Presentation, draft: bool) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_base(slide, "Project objective", "A prediction API can become a free training set", 2, draft)
    add_text(
        slide,
        "Modern pay-per-query APIs can become labeling services: attackers collect input-output pairs to train a substitute without the weights or original training data.",
        0.55,
        1.37,
        11.8,
        0.62,
        11,
        INK,
    )
    add_flow(
        slide,
        (
            ("1. Query", "Submit selected or manipulated images"),
            ("2. Collect", "Record labels and confidence values"),
            ("3. Train", "Use API answers as surrogate labels"),
            ("4. Replace", "Imitate the protected model's behavior"),
        ),
        2.30,
    )
    add_panel(slide, 0.55, 4.05, 3.62, 1.65)
    add_text(slide, "BUYER / USER", 0.77, 4.31, 2.8, 0.25, 8, TEAL, True)
    add_text(slide, "Government and commercial AI API owners; SOC and API-security teams", 0.77, 4.72, 3.0, 0.66, 10.5, INK, True)
    add_panel(slide, 4.36, 4.05, 3.62, 1.65)
    add_text(slide, "BUSINESS OUTCOME", 4.58, 4.31, 2.8, 0.25, 8, TEAL, True)
    add_text(slide, "Protect model IP and API revenue while preserving legitimate access", 4.58, 4.72, 3.0, 0.66, 11, INK, True)
    add_panel(slide, 8.17, 4.05, 4.29, 1.65, PANEL_LIGHT)
    add_text(slide, "THE GAP", 8.39, 4.31, 2.8, 0.25, 8, GOLD, True)
    add_text(slide, "Replaces manual query-log review with real-time behavioral containment; complements WAFs and rate limits.", 8.39, 4.72, 3.65, 0.70, 10.5, INK, True)
    set_notes(
        slide,
        "[Timing: 0:35-1:05]\nWalk left to right through the extraction loop. The attacker queries, collects labels, trains a different model, and can eventually replace paid API access. Explain the buyer and business impact. Emphasize the design gap: rate limiting is useful but cannot characterize the information value and structure of a query sequence. Transition: ModelSentry converts that sequence into persistent evidence.",
    )


def add_slide_three(prs: Presentation, draft: bool) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_base(slide, "Proposed solution", "Working prototype: stateful evidence, graduated response", 3, draft)
    add_flow(
        slide,
        (
            ("FastAPI", "Validated image and API-key context"),
            ("Victim CNN", "Label, confidence, margin, embedding"),
            ("50-query monitor", "Per-key rolling behavioral evidence"),
            ("Policy + dashboard", "Allow, observe, throttle, block; persist evidence"),
        ),
        1.48,
    )
    add_text(slide, "MULTI-TIMESCALE EVIDENCE", 0.55, 3.12, 4.0, 0.25, 8.5, TEAL, True)
    signals = (
        ("RATE", "Harvesting velocity"),
        ("DIVERSITY", "Embedding exploration"),
        ("REPLAY", "Exact request repetition"),
        ("BOUNDARY", "Low-margin probing"),
        ("LINKAGE", "Simulated account groups"),
    )
    for index, (heading, detail) in enumerate(signals):
        x = 0.55 + index * 2.40
        add_panel(slide, x, 3.55, 2.17, 1.08)
        add_text(slide, heading, x + 0.15, 3.75, 1.8, 0.22, 8.5, GOLD, True)
        add_text(slide, detail, x + 0.15, 4.10, 1.8, 0.30, 8, MUTED)
    add_panel(slide, 0.55, 5.03, 11.91, 1.22, PANEL_LIGHT)
    add_text(slide, "INFORMATION ACQUISITION PROXY", 0.78, 5.31, 3.2, 0.23, 8.5, TEAL, True)
    add_text(
        slide,
        "Short-window signals remain telemetry until confirmed by persistent, repeated, boundary-heavy, extreme-rate, or simulated account-group evidence.",
        4.15,
        5.25,
        7.85,
        0.65,
        10.5,
        INK,
    )
    set_notes(
        slide,
        "[Timing: 1:05-1:50]\nFollow one request through the system. The victim model returns the normal prediction while exposing internal telemetry to the monitor. Explain that V2.4 combines short-window model-aware telemetry with independent long-horizon confirmation, including exact replay and linked accounts. This separation is why none of 90 benign sessions was mitigated. Transition: we tested the frozen policy against real surrogate-training loops.",
    )


def add_slide_four(
    prs: Presentation,
    summary: dict,
    frame: pd.DataFrame,
    chart_path: Path,
    draft: bool,
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_base(slide, "Solution validation", "11/12 attacks detected; 0/90 benign sessions incorrectly mitigated", 4, draft)
    add_panel(slide, 0.55, 1.40, 7.45, 4.90)
    slide.shapes.add_picture(str(chart_path), Inches(0.79), Inches(1.72), width=Inches(6.97), height=Inches(3.92))
    add_text(slide, "Four modes evaluated on identical attacks and benign traffic", 0.86, 5.83, 6.8, 0.23, 8, MUTED, align=PP_ALIGN.CENTER)
    accuracy = frame.drop_duplicates("seed")["victim_validation_accuracy"]
    enhanced = summary["mode_overview"]["enhanced"]
    api_latency = summary["latency"]["enhanced/api_in_process"]["p50_ms"]["mean"]
    add_metric(
        slide,
        8.28,
        1.40,
        4.18,
        f"{accuracy.mean():.2%} +/- {accuracy.std() * 100:.2f} pp",
        "Victim validation accuracy",
        TEAL,
    )
    add_metric(slide, 8.28, 2.65, 4.18, "11/12 (91.7%)", "Enhanced attack detection")
    add_metric(
        slide,
        8.28,
        3.90,
        4.18,
        f"{enhanced['final_fidelity']['mean']:.2%} +/- {enhanced['final_fidelity']['std'] * 100:.2f} pp",
        "Mean final fidelity; lower is better",
        TEAL,
    )
    api_p95 = summary["latency"]["enhanced/api_in_process"]["p95_ms"]["mean"]
    add_metric(slide, 8.28, 5.15, 4.18, f"{api_latency:.2f} / {api_p95:.2f} ms", "API p50 / p95")
    add_text(slide, "Frozen 3-seed Fashion-MNIST holdout  |  simulated traffic  |  disjoint splits  |  different SGD surrogate  |  12 attack runs", 0.75, 6.53, 11.8, 0.25, 8, MUTED, align=PP_ALIGN.CENTER)
    set_notes(
        slide,
        "[Timing: 1:50-2:40]\nState the methodology before the result: three untouched holdout seeds, 40,000 victim-training images, eight epochs, four attack types, four detector modes, and 90 benign sessions per mode. Enhanced V2.4 detected 11 of 12 attack runs and mitigated none of the 90 benign sessions. Mean final surrogate fidelity was 71.78 percent. The development result did not reproduce perfectly: one slow-adaptive run was missed. Transition: the prototype is useful, but its limits are explicit.",
    )


def add_slide_five(prs: Presentation, dashboard_path: Path, qr_path: Path, video_url: str, draft: bool) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_base(slide, "Results and conclusions", "Operationally useful, honestly bounded", 5, draft)
    add_panel(slide, 0.55, 1.42, 6.15, 3.46)
    slide.shapes.add_picture(str(dashboard_path), Inches(0.74), Inches(1.63), width=Inches(5.77), height=Inches(3.24))
    add_panel(slide, 6.95, 1.42, 5.51, 3.46)
    takeaways = (
        ("01", "Validated improvement", "Same traffic: detection rose 4/12 to 11/12; benign mitigation fell 4/90 to 0/90."),
        ("02", "Lower attacker fidelity", "Mean final fidelity fell from 80.92% to 71.78% against baseline full."),
        ("03", "Known residual risk", "One patient slow-adaptive holdout run was not detected."),
    )
    for index, (number, heading, detail) in enumerate(takeaways):
        y = 1.72 + index * 0.94
        add_text(slide, number, 7.20, y, 0.50, 0.25, 9, GOLD, True)
        add_text(slide, heading, 7.82, y - 0.02, 3.95, 0.26, 10.5, INK, True)
        add_text(slide, detail, 7.82, y + 0.32, 4.05, 0.42, 8.2, MUTED)
    add_panel(slide, 0.55, 5.18, 8.62, 1.05, PANEL_LIGHT)
    add_text(slide, "DEPLOYMENT ROADMAP", 0.78, 5.41, 2.0, 0.23, 8, TEAL, True)
    add_text(slide, "Identity resolver  ->  streaming storage  ->  drift review  ->  bounded history  ->  analyst feedback", 2.67, 5.37, 6.14, 0.42, 9.2, INK)
    add_panel(slide, 9.42, 5.18, 3.04, 1.05)
    slide.shapes.add_picture(str(qr_path), Inches(9.61), Inches(5.31), width=Inches(0.78), height=Inches(0.78))
    add_text(slide, "SOURCE CODE", 10.56, 5.35, 1.53, 0.22, 8, TEAL, True)
    add_text(slide, "github.com/AT-14/ModelSentry", 10.56, 5.65, 1.55, 0.40, 6.8, TEAL)
    add_link_overlay(slide, 10.50, 5.58, 1.70, 0.48, REPOSITORY)
    video_text = "Demo video: link pending" if draft else "Demo video: open recording"
    add_text(slide, video_text, 0.74, 6.48, 5.3, 0.24, 8, RED if draft else TEAL, True)
    if not draft:
        add_link_overlay(slide, 0.68, 6.41, 2.25, 0.38, video_url)
    set_notes(
        slide,
        "[Timing: 2:40-3:25]\nClose with three points. ModelSentry provides explainable warning, graduated containment, and an honest measured boundary. Name the exact gap: one slow-adaptive run was missed even though all fast, replay, and distributed runs were detected. Briefly give the production roadmap. Invite judges to scan the repository QR code. Finish: ModelSentry turns model extraction from an invisible billing pattern into an observable security incident.\n\n[Backup demo order]\nShow the healthy API, legitimate interactive and batch clients, alert reasons, and the frozen V2.4 aggregate result. If live services fail, use the saved dashboard snapshot and comparison chart.",
    )


def add_pdf_links(pdf_path: Path, video_url: str, draft: bool) -> None:
    document = pymupdf.open(pdf_path)
    page = document[4]
    height = page.rect.height
    existing = {link.get("uri") for link in page.get_links()}
    if REPOSITORY not in existing:
        page.insert_link(
            {
                "kind": pymupdf.LINK_URI,
                "from": pymupdf.Rect(678, height - 167, 897, height - 91),
                "uri": REPOSITORY,
            }
        )
    if not draft and video_url not in existing:
        page.insert_link(
            {
                "kind": pymupdf.LINK_URI,
                "from": pymupdf.Rect(53, height - 74, 435, height - 48),
                "uri": video_url,
            }
        )
    temporary_path = pdf_path.with_suffix(".linked.pdf")
    document.save(temporary_path, garbage=4, deflate=True)
    document.close()
    temporary_path.replace(pdf_path)


def export_pdf(presentation_path: Path, pdf_path: Path, video_url: str, draft: bool) -> None:
    import win32com.client

    application = win32com.client.DispatchEx("PowerPoint.Application")
    application.Visible = True
    presentation = None
    try:
        presentation = application.Presentations.Open(
            str(presentation_path.resolve()),
            ReadOnly=True,
            Untitled=False,
            WithWindow=False,
        )
        presentation.SaveAs(str(pdf_path.resolve()), 32)
    finally:
        if presentation is not None:
            presentation.Close()
        application.Quit()
    add_pdf_links(pdf_path, video_url, draft)


def main() -> None:
    parser = argparse.ArgumentParser(description="Build ModelSentry PPTX and matching PDF")
    parser.add_argument("--video-url", default=VIDEO_LINK_PENDING)
    parser.add_argument("--skip-pdf", action="store_true")
    args = parser.parse_args()
    draft = args.video_url == VIDEO_LINK_PENDING
    if not draft:
        args.video_url = validate_public_url(args.video_url)

    summary_path = EVIDENCE / "validation_summary_v2.json"
    metrics_path = EVIDENCE / "per_mode_metrics_v2.csv"
    chart_path = SUBMISSION / "v2_holdout_comparison.png"
    dashboard_path = SUBMISSION / "dashboard_snapshot.png"
    if not all(path.exists() for path in (summary_path, metrics_path)):
        raise FileNotFoundError("Accepted V2.4 holdout evidence is missing")
    summary = json.loads(summary_path.read_text(encoding="utf-8"))
    frame = pd.read_csv(metrics_path)
    SUBMISSION.mkdir(parents=True, exist_ok=True)
    print("Building presentation assets", flush=True)
    make_holdout_chart(summary, chart_path)
    make_dashboard_snapshot(summary, dashboard_path)
    qr_path = SUBMISSION / "repository_qr.png"
    make_qr(REPOSITORY, qr_path)

    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_W)
    presentation.slide_height = Inches(SLIDE_H)
    presentation.core_properties.title = "ModelSentry: Detect Model-Theft Behavior and Limit API Leakage"
    presentation.core_properties.subject = "School of Cyber Defence 2026"
    presentation.core_properties.author = f"{TEAM}: {', '.join(MEMBERS)}"
    presentation.core_properties.keywords = "model extraction, API security, cybersecurity"
    presentation.core_properties.comments = ""
    presentation.core_properties.last_modified_by = TEAM

    add_slide_one(presentation, summary, draft)
    add_slide_two(presentation, draft)
    add_slide_three(presentation, draft)
    add_slide_four(presentation, summary, frame, chart_path, draft)
    add_slide_five(presentation, dashboard_path, qr_path, args.video_url, draft)
    print("Built five presentation slides", flush=True)

    suffix = "_DRAFT" if draft else ""
    presentation_path = SUBMISSION / f"ModelSentry_Presentation{suffix}.pptx"
    pdf_path = SUBMISSION / f"ModelSentry_Submission{suffix}.pdf"
    print(f"Saving {presentation_path}", flush=True)
    presentation.save(presentation_path)
    print(f"Wrote {presentation_path}")
    if not args.skip_pdf:
        export_pdf(presentation_path, pdf_path, args.video_url, draft)
        print(f"Wrote {pdf_path}")


if __name__ == "__main__":
    main()
